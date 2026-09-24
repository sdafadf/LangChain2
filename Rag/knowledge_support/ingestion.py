"""Uploaded file -> documents -> chunks -> embeddings -> existing Milvus collection."""
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

MAX_BYTES = 10 * 1024 * 1024
MAX_CHUNKS = 1000
IMPORT_VERSION = 'upload-v1-200-8'


@dataclass
class PreparedDocument:
    document_id: str
    source: str
    chunks: list
    warnings: list


def load_document(uploaded_file):
    """Accept Streamlit UploadedFile (name + getvalue), never a hardcoded path."""
    name = str(uploaded_file.name).replace('\\', '/').rsplit('/', 1)[-1]
    content = uploaded_file.getvalue()
    if not content:
        raise ValueError('文件为空，请选择有内容的文档。')
    if len(content) > MAX_BYTES:
        raise ValueError('单个文件不能超过 10 MB。')
    suffix = Path(name).suffix.lower()
    warnings, documents = [], []
    if suffix == '.txt':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError as error:
            raise ValueError('TXT 请先另存为 UTF-8 编码后上传。') from error
        if '\x00' in text:
            raise ValueError('TXT 含异常二进制字符，请确认文件格式和编码。')
        documents = [Document(page_content=text, metadata={'source': name})]
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('暂不支持加密 PDF，请解密后上传。')
        if len(reader.pages) > 300:
            raise ValueError('PDF 最多支持 300 页，请拆分文件后上传。')
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if text.strip():
                documents.append(Document(page_content=text, metadata={'source': name, 'page': i}))
            else:
                warnings.append(f'第 {i} 页没有可提取文字，已跳过。扫描页需先进行 OCR。')
    elif suffix == '.docx':
        from docx import Document as WordDocument
        with ZipFile(BytesIO(content)) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('Word 解压后内容过大，请拆分文档。')
        word = WordDocument(BytesIO(content))
        from docx.text.paragraph import Paragraph
        blocks = []
        for block in word.iter_inner_content():
            if isinstance(block, Paragraph):
                blocks.append(block.text)
            else:
                blocks.extend(' | '.join(cell.text for cell in row.cells) for row in block.rows)
        documents = [Document(page_content='\n'.join(blocks), metadata={'source': name})]
        warnings.append('Word 提取正文段落和表格；图片、文本框和页眉页脚不在本次提取范围内。')
    else:
        raise ValueError('只支持 .txt、.pdf 和 .docx 文件。')
    if not any(doc.page_content.strip() for doc in documents):
        raise ValueError('没有提取到文字。扫描版文档需先进行 OCR，再上传可复制文字的版本。')
    if sum(len(doc.page_content) for doc in documents) > 200_000:
        raise ValueError('提取的文本超过 20 万字，请拆分文件。')
    return name, sha256(content).hexdigest(), documents, warnings


def split_documents(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=200, chunk_overlap=8,
        separators=['\n======================\n', '\n\n', '\n', '。', ' ', ''])
    chunks = splitter.split_documents(documents)
    chunks = [chunk for chunk in chunks if chunk.page_content.strip()]
    if len(chunks) > MAX_CHUNKS:
        raise ValueError('切分后超过 1000 个片段，请拆成更小的文件。')
    return chunks


def prepare_document(uploaded_file):
    source, document_id, documents, warnings = load_document(uploaded_file)
    return PreparedDocument(document_id, source, split_documents(documents), warnings)


def chunk_id(document_id, index):
    # Negative INT64 IDs leave the notebook's existing 0, 1, ... IDs untouched.
    digest = sha256(f'{IMPORT_VERSION}:{document_id}:{index}'.encode()).digest()
    return -(int.from_bytes(digest[:8], 'big') & ((1 << 63) - 1)) - 1


def save_to_milvus(prepared, kb, progress=None):
    """Idempotent import: matching content is skipped; failed batches can be retried."""
    report = progress or (lambda fraction, message: None)
    if not prepared.chunks:
        raise ValueError('没有可以导入的片段。')
    client, collection = kb.client, kb.settings.collection
    report(0.0, '检查集合结构和重复片段…')
    schema = client.describe_collection(collection, timeout=10)
    from pymilvus import DataType
    fields = {field['name']: field for field in schema['fields']}
    if (schema.get('auto_id') or not schema.get('enable_dynamic_field')
            or fields.get('id', {}).get('type') != DataType.INT64
            or not fields.get('id', {}).get('is_primary')
            or fields.get('vector', {}).get('type') != DataType.FLOAT_VECTOR):
        raise ValueError('上传需要 id 为手动 INT64 主键、vector 为稠密向量且开启动态字段的集合。')
    dimension = int(fields['vector']['params']['dim'])
    ids = [chunk_id(prepared.document_id, i) for i in range(len(prepared.chunks))]
    if len(ids) != len(set(ids)):
        raise ValueError('片段 ID 冲突，已停止导入。')
    existing = {}
    for offset in range(0, len(ids), 100):
        rows = client.get(collection_name=collection, ids=ids[offset:offset + 100],
                          output_fields=['document_id', 'chunk_id', 'embedding_model', 'import_version'],
                          consistency_level='Strong', timeout=15)
        existing.update({row['id']: row for row in rows})
    for i, row_id in enumerate(ids):
        row = existing.get(row_id)
        if row is not None and (row.get('document_id') != prepared.document_id
                                or row.get('chunk_id') != i
                                or row.get('embedding_model') != kb.settings.embedding
                                or row.get('import_version') != IMPORT_VERSION):
            raise ValueError('已有记录的 ID 或向量模型配置冲突，已停止导入，不覆盖原记录。')
    missing = [i for i, row_id in enumerate(ids) if row_id not in existing]
    if not missing:
        report(1.0, '该文件已完整导入，已跳过。')
        return {'status': 'duplicate', 'written': 0, 'total': len(ids)}
    from langchain.embeddings import init_embeddings
    written = 0
    with httpx.Client(trust_env=kb.settings.use_proxy, timeout=60) as transport:
        embedder = init_embeddings(model=kb.settings.embedding, request_timeout=60,
                                   max_retries=0, http_client=transport)
        for offset in range(0, len(missing), 32):
            batch = missing[offset:offset + 32]
            report(written / len(missing), f'生成向量：{written}/{len(missing)} 个待导入片段')
            vectors = embedder.embed_documents([prepared.chunks[i].page_content for i in batch])
            if len(vectors) != len(batch) or any(len(vector) != dimension for vector in vectors):
                raise ValueError(f'向量数量或维度不匹配，集合要求 {dimension} 维。请使用与原入库相同的模型。')
            data = []
            for i, vector in zip(batch, vectors):
                chunk = prepared.chunks[i]
                data.append({'id': ids[i], 'vector': vector, 'text': chunk.page_content,
                             'source': prepared.source, 'chunk_id': i,
                             'document_id': prepared.document_id, 'content_hash': prepared.document_id,
                             'embedding_model': kb.settings.embedding, 'import_version': IMPORT_VERSION,
                             'page': chunk.metadata.get('page'), 'chunk_count': len(ids)})
            report(written / len(missing), f'写入知识库：{written}/{len(missing)}')
            client.upsert(collection_name=collection, data=data, timeout=30)
            written += len(batch)
    client.flush(collection_name=collection, timeout=30)
    report(1.0, f'已导入 {written} 个片段')
    return {'status': 'imported', 'written': written, 'total': len(ids)}


def import_document(uploaded_file, kb, progress=None):
    """Convenience entry point for use from a script or notebook."""
    return save_to_milvus(prepare_document(uploaded_file), kb, progress)
