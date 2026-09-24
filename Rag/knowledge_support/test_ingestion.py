import unittest
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from langchain_core.documents import Document
from pymilvus import DataType

from backend import Settings
from ingestion import PreparedDocument, chunk_id, prepare_document, save_to_milvus


def upload(name, data):
    return SimpleNamespace(name=name, getvalue=lambda: data)


class FakeClient:
    def __init__(self):
        self.rows = {0: {'id': 0, 'text': '原有知识'}}
        self.calls = 0
        self.fail_on = None

    def describe_collection(self, *args, **kwargs):
        return {'auto_id': False, 'enable_dynamic_field': True, 'fields': [
            {'name': 'id', 'type': DataType.INT64, 'is_primary': True},
            {'name': 'vector', 'type': DataType.FLOAT_VECTOR, 'params': {'dim': 3}}]}

    def get(self, ids, **kwargs):
        return [self.rows[i] for i in ids if i in self.rows]

    def upsert(self, data, **kwargs):
        self.calls += 1
        if self.calls == self.fail_on:
            raise ConnectionError('写入中断')
        for row in data:
            self.rows[row['id']] = row

    def flush(self, **kwargs):
        pass


class IngestionTests(unittest.TestCase):
    def setUp(self):
        self.client = FakeClient()
        self.kb = SimpleNamespace(client=self.client, settings=Settings())
        self.embedder = Mock()
        self.embedder.embed_documents.side_effect = lambda texts: [[0.1, 0.2, 0.3] for _ in texts]
        self.factory = patch('langchain.embeddings.init_embeddings', return_value=self.embedder).start()
        self.addCleanup(patch.stopall)

    def test_txt_bom_and_stable_hash(self):
        first = prepare_document(upload('a.txt', '知识库退款条件。'.encode('utf-8-sig')))
        second = prepare_document(upload('renamed.txt', '知识库退款条件。'.encode('utf-8-sig')))
        self.assertEqual(first.document_id, second.document_id)
        self.assertEqual(first.chunks[0].page_content, '知识库退款条件。')
        self.assertFalse(self.factory.called)

    def test_reject_invalid_and_blank(self):
        for name, data in [('a.exe', b'abc'), ('a.txt', b''), ('a.txt', b'  \n'),
                           ('a.txt', b'\xff'), ('a.txt', b'\x00'), ('a.pdf', b'invalid')]:
            with self.subTest(name=name, data=data), self.assertRaises(Exception):
                prepare_document(upload(name, data))

    def test_docx_paragraph_and_table(self):
        from docx import Document as Word
        document = Word()
        document.add_paragraph('服务时间')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '工作日'
        table.cell(0, 1).text = '9点至18点'
        buffer = BytesIO()
        document.save(buffer)
        prepared = prepare_document(upload('service.docx', buffer.getvalue()))
        text = '\n'.join(chunk.page_content for chunk in prepared.chunks)
        self.assertIn('服务时间', text)
        self.assertIn('工作日 | 9点至18点', text)

    def test_pdf_empty_page_rejected(self):
        from pypdf import PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, 'OCR'):
            prepare_document(upload('scan.pdf', buffer.getvalue()))

    def test_upload_deduplicate_and_keep_old_rows(self):
        prepared = prepare_document(upload('a.txt', '第一份知识'.encode()))
        result = save_to_milvus(prepared, self.kb)
        self.assertEqual(result['written'], 1)
        self.assertEqual(self.client.rows[0]['text'], '原有知识')
        calls = self.embedder.embed_documents.call_count
        self.assertEqual(save_to_milvus(prepared, self.kb)['status'], 'duplicate')
        self.assertEqual(self.embedder.embed_documents.call_count, calls)
        other = prepare_document(upload('a.txt', '同名不同内容'.encode()))
        save_to_milvus(other, self.kb)
        self.assertEqual(len(self.client.rows), 3)
        self.assertTrue(all(i < 0 for i in self.client.rows if i != 0))

    def test_dimension_error_does_not_write(self):
        self.embedder.embed_documents.side_effect = lambda texts: [[1.0] for _ in texts]
        with self.assertRaisesRegex(ValueError, '维度'):
            save_to_milvus(prepare_document(upload('a.txt', b'hello')), self.kb)
        self.assertEqual(len(self.client.rows), 1)

    def test_partial_failure_retry_only_missing(self):
        prepared = PreparedDocument('a' * 64, 'a.txt', [Document(page_content=f'知识 {i}') for i in range(35)], [])
        self.client.fail_on = 2
        with self.assertRaises(ConnectionError):
            save_to_milvus(prepared, self.kb)
        self.assertEqual(len(self.client.rows), 33)
        self.client.fail_on = None
        result = save_to_milvus(prepared, self.kb)
        self.assertEqual(result['written'], 3)
        self.assertEqual(len(self.client.rows), 36)

    def test_collision_does_not_overwrite(self):
        prepared = prepare_document(upload('a.txt', b'hello'))
        row_id = chunk_id(prepared.document_id, 0)
        self.client.rows[row_id] = {'id': row_id, 'document_id': 'different'}
        with self.assertRaisesRegex(ValueError, '冲突'):
            save_to_milvus(prepared, self.kb)
        self.assertEqual(self.client.rows[row_id]['document_id'], 'different')
        self.assertFalse(self.factory.called)


if __name__ == '__main__':
    unittest.main()
